from __future__ import annotations

import hashlib
import re
import tomllib
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def test_phase3_configuration_preserves_approved_boundaries() -> None:
    with (ROOT / "config" / "ingestion.phase3.toml").open("rb") as stream:
        config = tomllib.load(stream)
    assert config["bronze"]["immutable"] is True
    assert config["fabric"]["audit_table"] == "ingestion_audit"
    assert config["file_transport"]["preferred"] == "sharepoint_online_file"
    assert config["file_transport"]["fallback"] == "onelake_demo_staging"
    assert config["file_transport"]["selected"] == "onelake_demo_staging"
    assert config["file_transport"]["demonstrated_in_fabric"] is True
    assert config["atlas_erp"]["snapshot_mode"] == "full"
    assert config["maintcontrol"]["snapshot_mode"] == "full"
    assert config["maintcontrol"]["live_connection_status"] == (
        "accepted_quick_tunnel_demo_validated_in_fabric"
    )


def test_maintcontrol_demo_scripts_keep_runtime_secrets_external() -> None:
    api_script = (ROOT / "scripts" / "start-maintcontrol-api.ps1").read_text(encoding="utf-8")
    tunnel_script = (ROOT / "scripts" / "start-maintcontrol-quick-tunnel.ps1").read_text(
        encoding="utf-8"
    )
    smoke_script = (ROOT / "scripts" / "test-maintcontrol-endpoints.ps1").read_text(
        encoding="utf-8"
    )
    combined = api_script + tunnel_script + smoke_script

    assert 'BindHost = "127.0.0.1"' in api_script
    assert 'Origin = "http://127.0.0.1:8001"' in tunnel_script
    assert "MAINTCONTROL_API_TOKEN" in combined
    assert re.search(r"https://[a-z0-9-]+\.trycloudflare\.com", combined) is None


def test_six_pipeline_names_and_batched_notebook_pattern_are_versioned() -> None:
    specification = (ROOT / "fabric" / "pipeline-build-spec.md").read_text(encoding="utf-8")
    expected = {
        "pl_ingest_atlas_erp",
        "pl_ingest_mes",
        "pl_ingest_quality",
        "pl_ingest_maintcontrol",
        "pl_ingest_technical_documents",
        "pl_ingest_all_sources",
    }
    assert all(f"## `{name}`" in specification for name in expected)
    assert "at most two invocations" in specification
    assert "hand-written pipeline export JSON" not in specification


def test_audit_notebook_has_exact_phase3_statuses_and_fields() -> None:
    notebook = (ROOT / "fabric" / "notebooks" / "nb_bronze_ingestion_audit.py").read_text(
        encoding="utf-8"
    )
    for status in (
        "STARTED",
        "SUCCEEDED",
        "FAILED",
        "SKIPPED_ALREADY_INGESTED",
        "CONFLICT_SOURCE_CHANGED",
        "SUCCEEDED_REPLAY",
    ):
        assert status in notebook
    for field in (
        "source_identity_key",
        "transport_source",
        "content_sha256",
        "force_reprocess",
        "replay_of_batch_id",
        "extract_window_start",
        "extract_window_end",
    ):
        assert f'"{field}"' in notebook


def test_quality_and_pdf_fixtures_are_binary_ready() -> None:
    workbook_path = ROOT / "sample-data" / "quality" / "quality_department_2025.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Inspections", "DefectTypes", "Targets"]
    workbook.close()

    pdfs = sorted((ROOT / "sample-data" / "technical_documents").glob("*.pdf"))
    assert len(pdfs) == 6
    hashes = {_sha256(path) for path in pdfs}
    assert len(hashes) == 6
    assert all(path.read_bytes().startswith(b"%PDF-") for path in pdfs)
