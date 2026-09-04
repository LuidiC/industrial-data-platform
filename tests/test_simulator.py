from __future__ import annotations

import csv
import json
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path

import yaml
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from atlas_simulator.api import create_app
from atlas_simulator.config import load_config
from atlas_simulator.generator import RULE_EXPECTATIONS, generate, quality_status_rule, shift_for
from atlas_simulator.writers import document_inputs, quarter_rows, write_all, write_pdfs

ROOT = Path(__file__).resolve().parents[1]


def config(tmp_path: Path):
    return load_config(ROOT / "config" / "simulator.default.toml", tmp_path / "output")


def monthly_config(tmp_path: Path, month: int):
    base = config(tmp_path)
    end = date(2025 + (month == 12), month % 12 + 1, 1) - date.resolution
    return replace(base, start_date=date(2025, month, 1), end_date=end)


def test_generation_is_deterministic_and_period_limited(tmp_path: Path) -> None:
    january = generate(monthly_config(tmp_path, 1), inject_anomalies=False)
    repeated = generate(monthly_config(tmp_path, 1), inject_anomalies=False)
    assert january.production_events == repeated.production_events
    assert len(january.production_orders) == 30
    assert all(row["event_started_at"].startswith("2025-01") for row in january.production_events)
    assert all(row["planned_start"].startswith("2025-01") for row in january.work_orders)


def test_independent_monthly_generation_matches_annual_business_data(tmp_path: Path) -> None:
    annual = generate(config(tmp_path), inject_anomalies=False)
    annual_events = {row["event_id"]: row for row in annual.production_events}
    annual_orders = {row["production_order_id"]: row for row in annual.production_orders}
    for month in range(1, 13):
        independent = generate(monthly_config(tmp_path, month), inject_anomalies=False)
        assert {row["event_id"]: row for row in independent.production_events} == {
            key: row
            for key, row in annual_events.items()
            if row["event_started_at"].startswith(f"2025-{month:02d}")
        }
        assert {row["production_order_id"]: row for row in independent.production_orders} == {
            key: row
            for key, row in annual_orders.items()
            if row["scheduled_start"].startswith(f"2025-{month:02d}")
        }


def test_clean_machine_intervals_never_overlap(tmp_path: Path) -> None:
    data = generate(config(tmp_path), inject_anomalies=False)
    for machine_id in {row["machine_id"] for row in data.production_events}:
        rows = sorted(
            (row for row in data.production_events if row["machine_id"] == machine_id),
            key=lambda row: row["event_started_at"],
        )
        assert all(
            datetime.fromisoformat(left["event_ended_at"])
            <= datetime.fromisoformat(right["event_started_at"])
            for left, right in zip(rows, rows[1:], strict=False)
        )


def test_clean_production_does_not_overlap_maintenance_and_lifecycle_is_coherent(
    tmp_path: Path,
) -> None:
    data = generate(config(tmp_path), inject_anomalies=False)
    work_orders = {row["work_order_id"]: row for row in data.work_orders}
    for event in data.maintenance_events:
        work_order = work_orders[event["work_order_id"]]
        assert datetime.fromisoformat(work_order["opened_at"]) <= datetime.fromisoformat(
            event["started_at"]
        )
        assert datetime.fromisoformat(work_order["closed_at"]) >= datetime.fromisoformat(
            event["ended_at"]
        )
        assert work_order["status"] == "completed"
    for event in data.production_events:
        start, end = (
            datetime.fromisoformat(event["event_started_at"]),
            datetime.fromisoformat(event["event_ended_at"]),
        )
        assert not any(
            maintenance["machine_id"] == event["machine_id"]
            and start < datetime.fromisoformat(maintenance["ended_at"])
            and end > datetime.fromisoformat(maintenance["started_at"])
            for maintenance in data.maintenance_events
        )


def test_shift_boundaries_cover_midnight_and_month_edges() -> None:
    assert shift_for(datetime.fromisoformat("2025-01-01T00:00:00-03:00")) == (
        "SHIFT-C",
        date(2024, 12, 31),
    )
    assert shift_for(datetime.fromisoformat("2025-01-01T05:59:00-03:00")) == (
        "SHIFT-C",
        date(2024, 12, 31),
    )
    assert shift_for(datetime.fromisoformat("2025-01-01T06:00:00-03:00"))[0] == "SHIFT-A"
    assert shift_for(datetime.fromisoformat("2025-01-31T22:00:00-03:00")) == (
        "SHIFT-C",
        date(2025, 1, 31),
    )
    assert shift_for(datetime.fromisoformat("2025-02-01T00:00:00-03:00")) == (
        "SHIFT-C",
        date(2025, 1, 31),
    )


def test_quality_status_semantics_are_disjoint() -> None:
    assert quality_status_rule("PASS") is None
    assert quality_status_rule("pass") == "DQ-QUAL-004"
    assert quality_status_rule("PENDING") == "DQ-QUAL-002"


def test_default_manifest_covers_every_documented_injected_rule(tmp_path: Path) -> None:
    data = generate(config(tmp_path))
    assert RULE_EXPECTATIONS.keys() <= {entry["rule_id"] for entry in data.anomaly_manifest}


def test_manifest_matches_serialized_artifacts_and_monthly_equivalence(tmp_path: Path) -> None:
    annual_config = config(tmp_path)
    annual = generate(annual_config)
    write_all(annual, annual_config, root=tmp_path / "annual", include_pdfs=False)
    manifest = json.loads(
        (tmp_path / "annual" / "metadata" / "anomaly_manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["effective_period"] == "2025-01-01/2025-12-31"
    assert manifest["entries"]
    for entry in manifest["entries"]:
        artifact = tmp_path / "annual" / entry["physical_file"]
        assert artifact.exists()
        if entry["format"] == "CSV":
            with artifact.open(newline="", encoding="utf-8") as stream:
                rows = list(csv.DictReader(stream))
            assert rows[entry["serialized_location"]["row_number"] - 2][
                entry["affected_field"]
            ] == str(entry["injected_value"])
        elif entry["format"] == "XLSX":
            sheet = load_workbook(artifact, read_only=True)[entry["worksheet"]]
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            row = list(
                sheet.iter_rows(
                    min_row=entry["serialized_location"]["row_number"],
                    max_row=entry["serialized_location"]["row_number"],
                    values_only=True,
                )
            )[0]
            actual = row[headers.index(entry["affected_field"])]
            assert ("" if actual is None else str(actual)) == str(entry["injected_value"])
        else:
            rows = json.loads(artifact.read_text(encoding="utf-8"))
            index = int(entry["serialized_location"]["json_pointer"].strip("/"))
            assert rows[index][entry["affected_field"]] == entry["injected_value"]
    ignored = {
        "physical_file",
        "format",
        "worksheet",
        "serialized_location",
        "source_row_index",
        "duplicate_of_source_row_index",
    }
    for month in range(1, 13):
        independent_config = monthly_config(tmp_path, month)
        independent = generate(independent_config)
        write_all(
            independent, independent_config, root=tmp_path / f"month-{month}", include_pdfs=False
        )
        independent_manifest = json.loads(
            (tmp_path / f"month-{month}" / "metadata" / "anomaly_manifest.json").read_text(
                encoding="utf-8"
            )
        )
        annual_entries = [
            entry
            for entry in manifest["entries"]
            if entry["effective_period"] == f"2025-{month:02d}"
        ]
        assert [
            {key: value for key, value in entry.items() if key not in ignored}
            for entry in annual_entries
        ] == [
            {key: value for key, value in entry.items() if key not in ignored}
            for entry in independent_manifest["entries"]
        ]


def _contracts() -> dict[str, dict]:
    contracts = {}
    for path in (ROOT / "docs" / "data-contracts").glob("*.yaml"):
        if path.name != "contract-template.yaml":
            contract = yaml.safe_load(path.read_text(encoding="utf-8"))
            contracts[contract["dataset"]] = contract
    return contracts


def _assert_contract_type(value: object, declared_type: str) -> None:
    if declared_type == "string":
        assert isinstance(value, str)
    elif declared_type == "integer":
        assert isinstance(value, int) and not isinstance(value, bool)
    elif declared_type == "boolean":
        assert isinstance(value, bool)
    elif declared_type == "decimal":
        assert isinstance(value, (int, float)) and not isinstance(value, bool)
    elif declared_type == "date":
        assert isinstance(value, str)
        date.fromisoformat(value)
    elif declared_type == "timestamp":
        assert isinstance(value, str)
        assert datetime.fromisoformat(value).tzinfo is not None
    else:
        raise AssertionError(f"Unsupported contract type: {declared_type}")


def _assert_clean_dq_semantics(data) -> None:
    assert all(row["quantity_produced"] >= 0 for row in data.production_events)
    assert all(row["status"] == "completed" for row in data.production_events)
    assert all(row["event_id"].startswith("PEV-") for row in data.production_events)
    assert all(row["inspection_id"] for row in data.inspections)
    assert all(quality_status_rule(row["quality_status"]) is None for row in data.inspections)
    assert all(row["defective_quantity"] <= row["inspected_quantity"] for row in data.inspections)
    assert all(
        datetime.fromisoformat(row["ended_at"]) >= datetime.fromisoformat(row["started_at"])
        and row["downtime_minutes"] >= 0
        for row in data.maintenance_events
    )
    assert all(
        row["maintenance_type"] != "corrective"
        or row["status"] != "completed"
        or row["failure_classification"]
        for row in data.work_orders
    )


def test_contracts_validate_clean_baseline_semantics(tmp_path: Path) -> None:
    data = generate(config(tmp_path), inject_anomalies=False)
    contracts = _contracts()
    rule_ids = {
        rule["id"]
        for rule in yaml.safe_load((ROOT / "docs" / "data-quality" / "rules.yaml").read_text())[
            "rules"
        ]
    }
    assert set(contracts) == set(data.source_rows())
    for dataset, rows in data.source_rows().items():
        fields = contracts[dataset]["fields"]
        assert [field["name"] for field in fields] == list(rows[0])
        for field in fields:
            name = field["name"]
            values = [row[name] for row in rows]
            if field["required"]:
                assert all(value not in (None, "") for value in values)
            for value in values:
                if value is not None:
                    _assert_contract_type(value, field["type"])
            if field["unique"]:
                assert len(values) == len(set(values))
            allowed = field["constraints"].get(
                "allowed_values", field["constraints"].get("canonical_allowed_values")
            )
            if allowed:
                assert set(values) <= set(allowed)
            assert set(field["related_dq_rules"]) <= rule_ids
    _assert_clean_dq_semantics(data)


def test_duplicate_anomalies_have_pristine_source_lineage(tmp_path: Path) -> None:
    generated = generate(config(tmp_path))
    generated_config = config(tmp_path)
    write_all(generated, generated_config, root=tmp_path / "output", include_pdfs=False)
    manifest = json.loads(
        (tmp_path / "output" / "metadata" / "anomaly_manifest.json").read_text(encoding="utf-8")
    )
    for rule_id, dataset, field in (
        ("DQ-PROD-006", "production_events", "event_id"),
        ("DQ-QUAL-005", "inspections", "inspection_id"),
    ):
        entries = [entry for entry in manifest["entries"] if entry["rule_id"] == rule_id]
        assert entries
        source_rows = getattr(generated, dataset)
        non_duplicate_indexes = {
            entry["source_row_index"]
            for entry in generated.anomaly_manifest
            if entry["dataset"] == dataset and entry["rule_id"] != rule_id
        }
        for entry in entries:
            assert entry["stable_pre_mutation_locator"]
            assert entry["source_record_locator"] == entry["stable_pre_mutation_locator"]
            assert entry["duplicate_record_locator"]
            assert entry["source_record_locator"] != entry["duplicate_record_locator"]
            assert entry["injected_value"] not in {"", "bad-event"}
            assert entry["duplicate_of_source_row_index"] not in non_duplicate_indexes
            source = source_rows[entry["duplicate_of_source_row_index"]]
            duplicate = source_rows[entry["source_row_index"]]
            assert source == duplicate
            assert entry["source_record_locator"] == (
                f"{dataset}:{entry['effective_period']}:source:{source[field]}"
            )
            assert entry["duplicate_record_locator"] == (
                f"{dataset}:{entry['effective_period']}:duplicate:{duplicate[field]}"
            )
            artifact = tmp_path / "output" / entry["physical_file"]
            if entry["format"] == "CSV":
                with artifact.open(newline="", encoding="utf-8") as stream:
                    serialized = list(csv.DictReader(stream))
                assert sum(row[field] == duplicate[field] for row in serialized) == 2
            else:
                sheet = load_workbook(artifact, read_only=True)[entry["worksheet"]]
                headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
                values = [
                    row[headers.index(field)]
                    for row in sheet.iter_rows(min_row=2, values_only=True)
                ]
                assert values.count(duplicate[field]) == 2


def test_committed_sample_manifest_covers_quality_anomalies() -> None:
    manifest = json.loads(
        (ROOT / "sample-data" / "metadata" / "anomaly_manifest.json").read_text(encoding="utf-8")
    )
    assert {"DQ-QUAL-002", "DQ-QUAL-003", "DQ-QUAL-004"} <= {
        entry["rule_id"] for entry in manifest["entries"]
    }


def test_xlsx_is_semantically_deterministic_and_pdf_uses_calendar_quarters(tmp_path: Path) -> None:
    data = generate(config(tmp_path))
    write_all(data, config(tmp_path), root=tmp_path / "first", include_pdfs=False)
    write_all(data, config(tmp_path), root=tmp_path / "second", include_pdfs=False)

    def cells(path: Path, title: str) -> list[tuple]:
        return list(load_workbook(path, read_only=True)[title].iter_rows(values_only=True))

    first = tmp_path / "first" / "quality" / "quality_department_2025.xlsx"
    second = tmp_path / "second" / "quality" / "quality_department_2025.xlsx"
    assert [cells(first, title) for title in ("Inspections", "DefectTypes", "Targets")] == [
        cells(second, title) for title in ("Inspections", "DefectTypes", "Targets")
    ]
    for quarter in range(1, 4):
        inputs = document_inputs(data, quarter)
        maintenance = quarter_rows(data.maintenance_events, "started_at", quarter)
        inspections = quarter_rows(data.inspections, "inspection_timestamp", quarter)
        assert inputs["maintenance"] == maintenance
        assert inputs["quality"] == inspections
        assert len(inputs["maintenance"]) == len(maintenance)
        assert len(inputs["quality"]) == len(inspections)
        assert all(
            (datetime.fromisoformat(row["started_at"]).month - 1) // 3 + 1 == quarter
            for row in inputs["maintenance"]
        )
        assert all(
            (datetime.fromisoformat(row["inspection_timestamp"]).month - 1) // 3 + 1 == quarter
            for row in inputs["quality"]
        )
    assert len(write_pdfs(data, tmp_path / "pdfs")) == 6


def test_api_validates_bad_client_inputs(tmp_path: Path) -> None:
    client = TestClient(create_app(generate(config(tmp_path))))
    headers = {"Authorization": "Bearer local-development-token"}
    assert client.get("/api/v1/work-orders?page_size=2", headers=headers).status_code == 200
    for url in (
        "/api/v1/work-orders?cursor=broken",
        "/api/v1/work-orders?cursor=LTE=",
        "/api/v1/work-orders?page_size=0",
        "/api/v1/work-orders?machine_id=bad",
        "/api/v1/maintenance-events?occurred_from=2025-01-01T00:00:00",
        "/api/v1/maintenance-events?occurred_from=2025-02-01T00:00:00-03:00&occurred_to=2025-01-01T00:00:00-03:00",
    ):
        response = client.get(url, headers=headers)
        assert 400 <= response.status_code < 500


def test_api_returns_empty_data_for_valid_unknown_machine(tmp_path: Path) -> None:
    client = TestClient(create_app(generate(config(tmp_path))))
    headers = {"Authorization": "Bearer local-development-token"}
    for endpoint in ("work-orders", "maintenance-events"):
        response = client.get(f"/api/v1/{endpoint}?machine_id=MCH-999", headers=headers)
        assert response.status_code == 200
        assert response.json()["data"] == []


def test_api_keeps_anomalous_maintenance_rows_in_unfiltered_payload(tmp_path: Path) -> None:
    client = TestClient(create_app(generate(config(tmp_path))))
    response = client.get(
        "/api/v1/maintenance-events?page_size=500",
        headers={"Authorization": "Bearer local-development-token"},
    )
    assert response.status_code == 200
    assert any(row["machine_id"] == "MCH-999" for row in response.json()["data"])
