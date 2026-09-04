from __future__ import annotations

import csv
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen.canvas import Canvas

from .config import SimulationConfig
from .generator import GENERATOR_VERSION, GeneratedData


def _write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_csv(data: GeneratedData, root: Path) -> list[Path]:
    written: list[Path] = []
    monthly: dict[str, list[dict]] = defaultdict(list)
    for row in data.production_events:
        monthly[row["event_started_at"][:7]].append(row)
    for period, rows in sorted(monthly.items()):
        path = root / "mes" / f"production_events_{period.replace('-', '_')}.csv"
        _write_csv(path, rows)
        written.append(path)
    return written


def write_quality_workbook(data: GeneratedData, root: Path) -> Path:
    path = root / "quality" / "quality_department_2025.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in (
        ("Inspections", data.inspections),
        ("DefectTypes", data.defect_types),
        ("Targets", data.targets),
    ):
        sheet = workbook.create_sheet(title)
        headers = list(rows[0])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in rows:
            sheet.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                28, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
    workbook.properties.created = datetime(2025, 1, 1)
    workbook.properties.modified = datetime(2025, 1, 1)
    workbook.save(path)
    return path


def write_api_fixtures(data: GeneratedData, root: Path) -> list[Path]:
    fixture_root = root / "maintcontrol"
    fixture_root.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, rows in (
        ("work_orders", data.work_orders),
        ("maintenance_events", data.maintenance_events),
    ):
        path = fixture_root / f"{name}.json"
        path.write_text(json.dumps(rows, indent=2, default=str), encoding="utf-8")
        paths.append(path)
    return paths


def write_manifest(data: GeneratedData, config: SimulationConfig, root: Path) -> Path:
    path = root / "metadata" / "anomaly_manifest.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    locations = {
        "production_events": ("mes", "CSV", None),
        "inspections": ("quality/quality_department_2025.xlsx", "XLSX", "Inspections"),
        "work_orders": ("maintcontrol/work_orders.json", "JSON", None),
        "maintenance_events": ("maintcontrol/maintenance_events.json", "JSON", None),
    }
    entries = []
    for entry in data.anomaly_manifest:
        relative, format_name, worksheet = locations[entry["dataset"]]
        physical_file = (
            f"mes/production_events_{entry['effective_period'].replace('-', '_')}.csv"
            if entry["dataset"] == "production_events"
            else relative
        )
        row_index = entry["source_row_index"]
        if entry["dataset"] == "production_events":
            row_index = sum(
                1
                for row in data.production_events[: entry["source_row_index"]]
                if row["event_started_at"].startswith(entry["effective_period"])
            )
        serialized_location = (
            {"row_number": row_index + 2}
            if format_name in {"CSV", "XLSX"}
            else {"json_pointer": f"/{entry['source_row_index']}"}
        )
        entries.append(
            entry
            | {
                "physical_file": physical_file,
                "format": format_name,
                "worksheet": worksheet,
                "serialized_location": serialized_location,
            }
        )
    document = {
        "generator_version": GENERATOR_VERSION,
        "configuration_schema_version": "1.0.0",
        "seed": config.seed,
        "effective_period": f"{config.start_date}/{config.end_date}",
        "source_rows": {name: len(rows) for name, rows in data.source_rows().items()},
        "entries": entries,
    }
    path.write_text(json.dumps(document, indent=2, default=str), encoding="utf-8")
    return path


def _pdf(canvas: Canvas, title: str, subtitle: str, lines: list[str]) -> None:
    width, height = letter
    canvas.setTitle(title)
    canvas.setFont("Helvetica-Bold", 16)
    canvas.drawString(0.75 * inch, height - 0.8 * inch, title)
    canvas.setFont("Helvetica", 10)
    canvas.drawString(0.75 * inch, height - 1.05 * inch, subtitle)
    canvas.line(0.75 * inch, height - 1.18 * inch, width - 0.75 * inch, height - 1.18 * inch)
    y = height - 1.55 * inch
    for line in lines:
        canvas.drawString(0.9 * inch, y, line)
        y -= 0.28 * inch
    canvas.setFont("Helvetica-Oblique", 8)
    canvas.drawString(
        0.75 * inch, 0.55 * inch, "Synthetic document - Atlas Industrial Manufacturing"
    )
    canvas.showPage()


def quarter_rows(rows: list[dict], timestamp_field: str, quarter: int) -> list[dict]:
    return [
        row
        for row in rows
        if (datetime.fromisoformat(row[timestamp_field]).month - 1) // 3 + 1 == quarter
    ]


def document_inputs(data: GeneratedData, quarter: int) -> dict[str, list[dict]]:
    """Return the exact source subsets used by the two lightweight PDF templates."""
    return {
        "maintenance": quarter_rows(data.maintenance_events, "started_at", quarter),
        "quality": quarter_rows(data.inspections, "inspection_timestamp", quarter),
    }


def write_pdfs(data: GeneratedData, root: Path) -> list[Path]:
    """Create exactly six one-page PDFs using only two intentionally simple templates."""
    document_root = root / "technical_documents"
    document_root.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for quarter in range(1, 4):
        inputs = document_inputs(data, quarter)
        maintenance = inputs["maintenance"]
        path = document_root / f"maintenance_report_2025_q{quarter}.pdf"
        canvas = Canvas(str(path), pagesize=letter, invariant=1)
        _pdf(
            canvas,
            "Maintenance Report",
            f"Q{quarter} 2025 - synthetic operational summary",
            [
                f"Events reviewed: {len(maintenance)}",
                "Scope: preventive and corrective maintenance",
                "Source: MaintControl",
            ],
        )
        canvas.save()
        paths.append(path)
    for quarter in range(1, 4):
        inspections = document_inputs(data, quarter)["quality"]
        path = document_root / f"quality_bulletin_2025_q{quarter}.pdf"
        canvas = Canvas(str(path), pagesize=letter, invariant=1)
        _pdf(
            canvas,
            "Quality Bulletin",
            f"Q{quarter} 2025 - synthetic departmental bulletin",
            [
                f"Inspections sampled: {len(inspections)}",
                "Scope: inspection outcomes and defect categories",
                "Source: Quality Department",
            ],
        )
        canvas.save()
        paths.append(path)
    return paths


def write_all(
    data: GeneratedData,
    config: SimulationConfig,
    root: Path | None = None,
    include_pdfs: bool = True,
) -> list[Path]:
    destination = root or config.output_root
    destination.mkdir(parents=True, exist_ok=True)
    paths = write_csv(data, destination)
    paths.append(write_quality_workbook(data, destination))
    paths.extend(write_api_fixtures(data, destination))
    paths.append(write_manifest(data, config, destination))
    if include_pdfs:
        paths.extend(write_pdfs(data, destination))
    return paths
